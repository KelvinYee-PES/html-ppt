"""Automated tests for Phase 2: Template Extraction, Past Deck Analysis, Store & Bundling."""
from __future__ import annotations

from pathlib import Path
import pytest
from pptx import Presentation

from htmlppt.brand.bundle import export_brand_bundle, import_brand_bundle
from htmlppt.brand.deck_analyzer import analyze_past_decks
from htmlppt.brand.model import BrandProfile
from htmlppt.brand.store import BrandStore
from htmlppt.brand.template_extractor import extract_brand_from_template
from htmlppt.analyze.segmenter import build_deck
from htmlppt.extract.renderer import snapshot_html
from htmlppt.layout.engine import layout_deck
from htmlppt.render.pptx_writer import write_pptx


@pytest.fixture
def sample_pptx(tmp_path):
    path = tmp_path / "sample_corporate.pptx"
    prs = Presentation()
    prs.slide_width = 12192000  # 16:9
    prs.slide_height = 6858000
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Acme Corp Deck"
    prs.save(str(path))
    return path


def test_template_extraction(sample_pptx):
    brand = extract_brand_from_template(sample_pptx, brand_name="Acme Corp")
    assert brand.name == "Acme Corp"
    assert brand.slide_size.width == 12192000
    assert brand.slide_size.height == 6858000
    assert brand.colors.accent1 is not None
    assert brand.fonts.heading.family != ""
    assert brand.template is not None
    assert len(brand.template.layouts) > 0


def test_deck_statistical_analysis(sample_pptx):
    brand = analyze_past_decks([sample_pptx])
    assert brand.fonts.heading.family != ""
    assert brand.fonts.sizes.title > 0
    assert len(brand.sources) == 1
    assert brand.sources[0].kind == "deck"


def test_brand_store(tmp_path):
    db_path = tmp_path / "test_brands.db"
    store = BrandStore(db_path=db_path)

    brand = BrandProfile(id="brand_test_1", organization="TestOrg", name="Alpha Brand")
    store.save_brand(brand)

    retrieved = store.get_brand("brand_test_1")
    assert retrieved is not None
    assert retrieved.name == "Alpha Brand"
    assert retrieved.organization == "TestOrg"

    brands = store.list_brands(organization="TestOrg")
    assert len(brands) == 1

    # File storage
    store.save_file("file_1", "logo.png", "image/png", b"fake_png_data")
    fname, ctype, fbytes = store.get_file("file_1")
    assert fname == "logo.png"
    assert fbytes == b"fake_png_data"

    # Delete
    assert store.delete_brand("brand_test_1") is True
    assert store.get_brand("brand_test_1") is None


def test_brand_bundle_export_import(tmp_path):
    db_path = tmp_path / "bundle_test.db"
    store = BrandStore(db_path=db_path)

    brand = BrandProfile(id="b_bundle", organization="Org", name="Bundle Brand")
    brand.logo = None
    store.save_brand(brand)

    # Export bundle
    bundle_bytes = export_brand_bundle(brand, store)
    assert len(bundle_bytes) > 100

    # Import into another store
    db_path2 = tmp_path / "bundle_test2.db"
    store2 = BrandStore(db_path=db_path2)
    imported = import_brand_bundle(bundle_bytes, store2)
    assert imported.id == "b_bundle"
    assert imported.name == "Bundle Brand"
    assert store2.get_brand("b_bundle") is not None


def test_template_mode_generation(sample_pptx, tmp_path):
    # Extract brand from template
    brand = extract_brand_from_template(sample_pptx, brand_name="Acme")
    
    # Run conversion using the template
    fixture_html = Path(__file__).parent / "fixtures" / "html" / "deck_slides.html"
    snap = snapshot_html(path=fixture_html)
    deck = build_deck(snap)
    deck_model = layout_deck(deck, brand)

    out_pptx = tmp_path / "templated_output.pptx"
    write_pptx(deck_model, output_path=out_pptx, template_path=sample_pptx)
    assert out_pptx.exists()

    # Re-open and verify
    prs = Presentation(str(out_pptx))
    assert len(prs.slides) == len(deck_model.slides)
