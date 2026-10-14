"""Automated tests for Phase 1: Core Engine, Layout, and Native PPTX Generation."""
from __future__ import annotations

import io
from pathlib import Path
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from htmlppt.analyze.segmenter import build_deck
from htmlppt.brand.presets import apply_preset, palette_brand
from htmlppt.extract.renderer import snapshot_html
from htmlppt.layout.engine import layout_deck
from htmlppt.render.pptx_writer import write_pptx

FIXTURES_DIR = Path(__file__).parent / "fixtures" / "html"


@pytest.fixture(scope="module")
def deck_snapshot():
    html_path = FIXTURES_DIR / "deck_slides.html"
    return snapshot_html(path=html_path)


def test_snapshot_extraction(deck_snapshot):
    assert deck_snapshot is not None
    assert "root" in deck_snapshot
    assert len(deck_snapshot.get("charts", [])) >= 1
    chart0 = deck_snapshot["charts"][0]["spec"]
    assert chart0["type"] == "bar"
    assert "DACH" in chart0["categories"]
    assert len(chart0["series"]) == 2


def test_semantic_segmentation(deck_snapshot):
    deck = build_deck(deck_snapshot)
    assert deck.source_kind == "deck"
    assert len(deck.slides) == 7
    # Slide 1 is title
    assert deck.slides[0].kind == "title"
    assert "Q3 2026 Business Review" in deck.slides[0].title
    # Slide 2 has KPIs
    assert any(b.type == "kpis" for b in deck.slides[1].blocks)
    # Slide 3 has Chart
    assert any(b.type == "chart" for b in deck.slides[2].blocks)
    # Slide 4 has multi-column
    assert any(b.column is not None for b in deck.slides[3].blocks)
    # Slide 5 has Table
    assert any(b.type == "table" for b in deck.slides[4].blocks)
    # Slide 6 has Cards/Process
    assert any(b.type == "cards" for b in deck.slides[5].blocks)
    # Slide 7 has Quote
    assert any(b.type == "quote" for b in deck.slides[6].blocks)


def test_layout_and_pptx_generation(deck_snapshot, tmp_path):
    deck = build_deck(deck_snapshot)
    brand = palette_brand("ocean")
    brand = apply_preset(brand, "corporate")

    deck_model = layout_deck(deck, brand)
    assert len(deck_model.slides) == 7

    out_pptx = tmp_path / "test_output.pptx"
    raw_bytes = write_pptx(deck_model, output_path=out_pptx)
    assert out_pptx.exists()
    assert len(raw_bytes) > 10000

    # Re-open presentation with python-pptx and verify native elements
    prs = Presentation(str(out_pptx))
    assert len(prs.slides) == 7

    # Slide 2: KPI slide must have shapes and text boxes (not pictures!)
    s2 = prs.slides[1]
    s2_types = [shape.shape_type for shape in s2.shapes]
    assert MSO_SHAPE_TYPE.PICTURE not in s2_types
    assert MSO_SHAPE_TYPE.AUTO_SHAPE in s2_types or MSO_SHAPE_TYPE.TEXT_BOX in s2_types

    # Slide 3: Chart slide must have a real native chart
    s3 = prs.slides[2]
    has_chart = any(shape.has_chart for shape in s3.shapes)
    assert has_chart, "Slide 3 must contain a native PowerPoint chart"
    chart_shape = next(shape for shape in s3.shapes if shape.has_chart)
    chart = chart_shape.chart
    assert len(chart.series) == 2

    # Slide 5: Table slide must have a real native table
    s5 = prs.slides[4]
    has_table = any(shape.has_table for shape in s5.shapes)
    assert has_table, "Slide 5 must contain a native PowerPoint table"
    table_shape = next(shape for shape in s5.shapes if shape.has_table)
    table = table_shape.table
    assert len(table.rows) == 5
    assert len(table.columns) == 4
    # Check text content of first header cell
    assert "Segment" in table.cell(0, 0).text_frame.text
