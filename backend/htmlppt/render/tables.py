"""Native PowerPoint table element rendering."""
from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide
from pptx.util import Emu, Pt

from htmlppt.brand.model import hex_to_rgb
from htmlppt.layout import model as lm
from htmlppt.render.text import ALIGN_MAP


def render_table_element(slide: Slide, el: lm.TableEl) -> None:
    """Render a native, editable PowerPoint table."""
    n_rows = len(el.rows)
    n_cols = len(el.col_widths)
    if not n_rows or not n_cols:
        return

    table_shape = slide.shapes.add_table(n_rows, n_cols, Emu(el.x), Emu(el.y), Emu(el.w), Emu(el.h))
    table = table_shape.table

    # Set column widths
    for ci, cw in enumerate(el.col_widths):
        table.columns[ci].width = Emu(cw)

    # Set row heights
    for ri, rh in enumerate(el.row_heights):
        table.rows[ri].height = Emu(rh)

    # Populate cells
    for ri, row_data in enumerate(el.rows):
        for ci, cell_data in enumerate(row_data):
            if ci >= n_cols:
                continue
            cell = table.cell(ri, ci)

            # Background fill
            if cell_data.fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(*hex_to_rgb(cell_data.fill))

            # Cell text
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_left = Emu(60000)
            tf.margin_right = Emu(60000)
            tf.margin_top = Emu(40000)
            tf.margin_bottom = Emu(40000)

            for pi, p_model in enumerate(cell_data.paras):
                p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                p.alignment = ALIGN_MAP.get(p_model.align, PP_ALIGN.LEFT)

                for r_model in p_model.runs:
                    run = p.add_run()
                    run.text = r_model.text
                    run.font.name = r_model.font
                    run.font.size = Pt(r_model.size)
                    run.font.bold = r_model.bold
                    run.font.italic = r_model.italic
                    if r_model.color:
                        run.font.color.rgb = RGBColor(*hex_to_rgb(r_model.color))
