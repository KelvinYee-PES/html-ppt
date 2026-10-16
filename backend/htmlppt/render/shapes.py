"""Native PowerPoint vector shapes, fills and borders."""
from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.slide import Slide
from pptx.util import Emu, Pt

from htmlppt.brand.model import hex_to_rgb
from htmlppt.layout import model as lm
from htmlppt.render.text import ALIGN_MAP

_A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

GEOM_MAP = {
    "rect": MSO_SHAPE.RECTANGLE,
    "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "line": MSO_SHAPE.RECTANGLE,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "chevron": MSO_SHAPE.CHEVRON,
    "homePlate": MSO_SHAPE.PENTAGON,
}


def _set_fill_alpha(shape, alpha: float) -> None:
    """Apply transparency to a solid-filled shape via <a:alpha> inside srgbClr."""
    if alpha >= 0.999:
        return
    spPr = shape._element.spPr
    solid = spPr.find(_qn("a:solidFill"))
    if solid is None:
        return
    srgb = solid.find(_qn("a:srgbClr"))
    if srgb is None:
        return
    srgb.append(parse_xml(f'<a:alpha {_A_NS} val="{int(max(0.0, min(1.0, alpha)) * 100000)}"/>'))


def _qn(tag: str):
    from pptx.oxml.ns import qn
    return qn(tag)


def _set_gradient(shape, colors: list[str], angle: float) -> None:
    """Replace solid fill with a linear gradient of 2+ stops."""
    if not colors:
        return
    spPr = shape._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill"):
        for el in spPr.findall(_qn(tag)):
            spPr.remove(el)
    n = len(colors)
    stops = "".join(
        f'<a:gs pos="{int(i / max(n - 1, 1) * 100000)}"><a:srgbClr val="{c.lstrip("#")}"/></a:gs>'
        for i, c in enumerate(colors)
    )
    ang = int(((angle - 90) % 360) * 60000)
    grad = parse_xml(
        f'<a:gradFill {_A_NS} rotWithShape="1"><a:gsLst>{stops}</a:gsLst>'
        f'<a:lin ang="{ang}" scaled="1"/></a:gradFill>'
    )
    prst = spPr.find(_qn("a:prstGeom"))
    if prst is not None:
        prst.addnext(grad)
    else:
        spPr.append(grad)


def _set_shadow(shape) -> None:
    spPr = shape._element.spPr
    for el in spPr.findall(_qn("a:effectLst")):
        spPr.remove(el)
    spPr.append(parse_xml(
        f'<a:effectLst {_A_NS}><a:outerShdw blurRad="90000" dist="38100" dir="5400000" rotWithShape="0">'
        f'<a:srgbClr val="000000"><a:alpha val="20000"/></a:srgbClr></a:outerShdw></a:effectLst>'
    ))


def render_shape_element(slide: Slide, el: lm.ShapeEl) -> None:
    """Render a native PowerPoint auto shape."""
    shape_type = GEOM_MAP.get(el.geom, MSO_SHAPE.RECTANGLE)
    shape = slide.shapes.add_shape(shape_type, Emu(el.x), Emu(el.y), Emu(el.w), Emu(el.h))

    # Corner radius adjustment for rounded rectangles
    if el.geom == "roundRect" and len(shape.adjustments) > 0:
        shape.adjustments[0] = min(max(el.radius, 0.0), 0.5)

    # Fill
    if el.gradient:
        _set_gradient(shape, el.gradient, el.gradient_angle)
    elif el.fill:
        shape.fill.solid()
        r, g, b = hex_to_rgb(el.fill)
        shape.fill.fore_color.rgb = RGBColor(r, g, b)
        _set_fill_alpha(shape, el.fill_alpha)
    else:
        shape.fill.background()

    # Border line
    if el.line and el.line_w > 0:
        shape.line.color.rgb = RGBColor(*hex_to_rgb(el.line))
        shape.line.width = Pt(el.line_w)
        if el.dash:
            shape.line.dash_style = MSO_LINE.DASH
    else:
        shape.line.fill.background()

    if el.shadow:
        _set_shadow(shape)

    # Shape text if present
    if el.paras:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Emu(el.inset)
        tf.margin_top = Emu(el.inset)
        tf.margin_right = Emu(el.inset)
        tf.margin_bottom = Emu(el.inset)

        for i, p_model in enumerate(el.paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN_MAP.get(p_model.align, PP_ALIGN.LEFT)
            for r_model in p_model.runs:
                run = p.add_run()
                run.text = r_model.text
                run.font.name = r_model.font
                run.font.size = Pt(r_model.size)
                run.font.bold = r_model.bold
                run.font.italic = r_model.italic
                if r_model.color:
                    run.font.color.rgb = RGBColor(*hex_to_rgb(r_model.color))
